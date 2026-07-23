"""
Generic lossless workbook capture (v3.0 P7, gap G-M2).

One serializer walks EVERY sheet of an OOXML workbook and dumps each
cell verbatim into a JSON-safe structure — the "render as is" store
(plan §2): capture is NOT gated by what any UI renders; the raw grids
survive label moves and future sheets. The v2.7 normalized parse
(``multicharts.py``) stays as the thin DERIVED aggregate on top.

Shape (``format: workbook.v1``)::

    {"format": "workbook.v1",
     "sheets": [{"name": "Strategy Analysis", "rows": [[cell, …], …]}, …]}

Cell encoding — JSON-safe, type- and SIGN-preserving (the MC gotchas:
signed-negative drawdown/PF cells, ``"$50"`` Point-Value strings):

- empty            → ``null``
- str / bool       → bare (bool checked BEFORE numbers — bool is an int
                     subclass)
- finite number, plain format → bare number (sign rides the JSON number)
- finite number, styled format → ``{"t": "n", "v": num, "f": fmt}`` —
  the number-format string is the plan's "ideally" capture; a ``%``
  format is the UI's percent hint
- NON-FINITE float → ``{"t": "n", "v": "nan" | "inf" | "-inf"}`` — never
  a bare token (F2 discipline: bare NaN/Infinity poisons every JSON door
  that later serves the capture)
- datetime / date / time → ``{"t": "dt", "v": iso}`` (+ ``"f"`` when a
  non-default format string exists)
- anything else    → ``{"t": "str", "v": str(x)}``

Row/sheet hygiene: trailing empty cells are trimmed per row and trailing
empty rows per sheet; INTERIOR blank rows/cells are preserved — they are
MultiCharts' section separators and the render layer keys on them.
Sheets that cannot be read (chartsheets, image-only Graphs sheets) come
back with ``rows: []`` — present, honestly empty.
"""
from __future__ import annotations

import logging
import math
from datetime import date, datetime, time
from io import BytesIO
from typing import Any, Dict, List

from core.backtest_adapters.base import BacktestAdapterError

log = logging.getLogger("backtest_adapters.workbook_capture")

# Fail-loud ceiling (no-silent-caps rule): a capture beyond this many
# cells is refused, not truncated. The real @ES export is ~10k cells;
# this bounds a pathological 20 MB workbook, not a real report.
MAX_CAPTURE_CELLS = 250_000

# number_format values that carry no information worth wrapping for.
_PLAIN_FORMATS = {"", "General", "@"}


def _encode_cell(value: Any, fmt: str) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):  # before int/float — bool subclasses int
        return value
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float)):
        f = float(value)
        if not math.isfinite(f):
            v = "nan" if math.isnan(f) else ("inf" if f > 0 else "-inf")
            return {"t": "n", "v": v}
        if fmt not in _PLAIN_FORMATS:
            return {"t": "n", "v": value, "f": fmt}
        return value
    if isinstance(value, (datetime, date, time)):
        out: Dict[str, Any] = {"t": "dt", "v": value.isoformat()}
        if fmt not in _PLAIN_FORMATS:
            out["f"] = fmt
        return out
    return {"t": "str", "v": str(value)}


def capture_workbook(file_bytes: bytes) -> Dict[str, Any]:
    """Losslessly capture every sheet of an OOXML workbook (see module doc)."""
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover — pinned in requirements
        raise BacktestAdapterError(f"openpyxl unavailable: {e}") from e
    if file_bytes[:2] != b"PK":
        raise BacktestAdapterError(
            "capture_workbook: not an OOXML workbook (missing PK magic)"
        )
    try:
        wb = openpyxl.load_workbook(
            BytesIO(file_bytes), read_only=True, data_only=True
        )
    except Exception as e:
        raise BacktestAdapterError(
            f"Could not open workbook for capture: {e}"
        ) from e
    try:
        sheets: List[Dict[str, Any]] = []
        cells = 0
        for name in wb.sheetnames:
            rows: List[List[Any]] = []
            try:
                ws = wb[name]
                for row in ws.iter_rows():
                    out_row: List[Any] = []
                    for cell in row:
                        # EmptyCell / chartsheet oddities: degrade per-cell,
                        # never lose the sheet.
                        try:
                            fmt = getattr(cell, "number_format", "General") or "General"
                            out_row.append(_encode_cell(cell.value, fmt))
                        except Exception:
                            out_row.append(None)
                    while out_row and out_row[-1] is None:  # trim trailing empties
                        out_row.pop()
                    cells += len(out_row)
                    if cells > MAX_CAPTURE_CELLS:
                        raise BacktestAdapterError(
                            f"Workbook too large to capture verbatim "
                            f"(> {MAX_CAPTURE_CELLS} cells) — refusing rather "
                            "than truncating silently."
                        )
                    rows.append(out_row)
            except BacktestAdapterError:
                raise
            except Exception as e:
                # Chartsheets (the MC "* Graphs" image sheets) and other
                # unreadable sheet types: present, honestly empty.
                log.info("[capture] sheet %r unreadable (%s) — captured empty",
                         name, type(e).__name__)
                rows = []
            while rows and not rows[-1]:  # trim trailing empty rows
                rows.pop()
            sheets.append({"name": str(name), "rows": rows})
        return {"format": "workbook.v1", "sheets": sheets}
    finally:
        wb.close()
