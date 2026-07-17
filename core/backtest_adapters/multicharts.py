"""
MultiCharts Strategy Performance Report adapter (v2.7 Phase 2, task 2.4).

Verified against a real ``@ES - 1 Minute`` export (2026-06-24 sample,
re-inspected 2026-07-16). Layout facts the parser is built on:

- The export is an OOXML workbook (``PK`` magic bytes) even when saved
  as ``.xml`` — sniff bytes, not the extension. 8 sheets; only three
  carry data: ``Strategy Analysis`` (summary; labels col A, All-Trades
  values col B), ``List of Trades`` (title row, blank row, HEADER ON
  ROW 3 — found by scanning, not hardcoded), ``Settings`` (key/value in
  cols A/B).
- Two rows per trade: Entry row (``Type``=Entry*, carries ``Trade #`` +
  ``Profit ($)``/``Cum. Profit ($)``/``Drawdown ($)``), then Exit row
  (``Type``=Exit*, carries ``Order #`` + exit ``Price`` + exit
  ``Signal``). A trailing unpaired Entry (position still open at export)
  is skipped.
- openpyxl (data_only) yields real ``datetime`` cells for dates; a
  numeric cell is still converted defensively via the Excel serial
  epoch (1899-12-30).
- Signed-value gotchas: ``Max Strategy Drawdown`` / ``(%)`` and the
  trade list's ``Drawdown ($)`` are NEGATIVE; ``Profit Factor`` is
  reported signed-negative (gross_profit / negative gross_loss). The
  normalized shape stores positive magnitudes and recomputes
  profit_factor = gross_profit / abs(gross_loss).
- ``Settings`` carries ``Point Value`` as a string like ``"$50"`` —
  size_usdt = entry Price x Contracts x Point Value (§6-3a); missing
  Point Value falls back to 1.0 (documented, logged).

Flat-XML note (named deviation from task 2.4's "ElementTree fallback"):
no true flat-XML/SpreadsheetML MC sample exists to validate a parser
against (§6-1: the MC ``.xml`` export IS an OOXML zip), so non-PK XML
is DETECTED via ElementTree and rejected with a specific error instead
of speculatively parsed. If a real flat-XML variant ever appears, add
the parse branch then — with its fixture.
"""
from __future__ import annotations

import logging
import math
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from core.backtest_adapters.base import BacktestAdapterError
from core.backtest_adapters.protocols import (
    BacktestSummary,
    BacktestTrade,
    EquityPoint,
    NormalizedBacktestResult,
)
from core.backtest_adapters.registry import register_backtest_adapter

log = logging.getLogger("backtest_adapters.multicharts")

_EXCEL_EPOCH = datetime(1899, 12, 30)


def _to_int(value: "float | None", fallback: int) -> int:
    """Non-finite-safe int coercion (v2.7 Task C, F7): None/nan/inf →
    fallback; everything else truncates. NB deliberate semantics change
    vs the old `int(x or fallback)`: a literal 0 cell now yields 0
    (faithful) instead of falling back."""
    if value is None or not math.isfinite(value):
        return fallback
    return int(value)


def _to_iso(value: Any) -> str:
    """Datetime cell -> ISO string; numeric cell -> via Excel serial epoch."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            return (_EXCEL_EPOCH + timedelta(days=float(value))).isoformat()
        except (OverflowError, ValueError):
            return ""
    return str(value)


def _to_float(value: Any) -> Optional[float]:
    """Tolerant numeric coercion: strips $ , and whitespace; None on failure.

    A trailing ``%`` converts to a FRACTION (``"37.2%"`` -> 0.372) — the
    normalized contract stores fractions, and silently keeping 37.2 would
    be a 100x scale drift (P2 audit MED-1; today's MC cells are numeric
    fractions, this guards a future format change).
    """
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("$", "").replace(",", "")
    if not s:
        return None
    is_pct = s.endswith("%")
    s = s.replace("%", "")
    if not s:
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return f / 100.0 if is_pct else f


@register_backtest_adapter("multicharts")
class MultiChartsAdapter:
    """MultiCharts Strategy Performance Report -> NormalizedBacktestResult."""

    app_id = "multicharts"
    display_name = "MultiCharts"
    accepted_extensions = (".xlsx", ".xml")

    def parse(self, file_bytes: bytes, filename: str) -> NormalizedBacktestResult:
        if file_bytes[:2] == b"PK":
            return self._parse_workbook(file_bytes)
        head = file_bytes.lstrip()[:200]
        if head.startswith(b"<"):
            # XML text, not an OOXML zip — detect (ET) and reject loudly.
            try:
                ET.fromstring(file_bytes.decode("utf-8", errors="replace"))
                kind = "flat-XML"
            except ET.ParseError:
                kind = "malformed XML"
            raise BacktestAdapterError(
                f"'{filename}' is {kind}, not the OOXML workbook MultiCharts "
                "exports (its .xml export is a zipped .xlsx). A true flat-XML "
                "variant is unsupported — no sample exists to validate against."
            )
        raise BacktestAdapterError(
            f"'{filename}' is not a MultiCharts Strategy Performance Report "
            "(unrecognized format — expected an OOXML workbook)."
        )

    # ── workbook ────────────────────────────────────────────────────────────

    def _parse_workbook(self, file_bytes: bytes) -> NormalizedBacktestResult:
        try:
            import openpyxl
        except ImportError as e:  # pragma: no cover — pinned in requirements
            raise BacktestAdapterError(f"openpyxl unavailable: {e}") from e
        try:
            wb = openpyxl.load_workbook(
                BytesIO(file_bytes), read_only=True, data_only=True
            )
        except Exception as e:
            raise BacktestAdapterError(
                f"Could not open workbook (corrupt or not an .xlsx): {e}"
            ) from e
        try:
            sheets = set(wb.sheetnames)
            if "List of Trades" not in sheets and "Strategy Analysis" not in sheets:
                raise BacktestAdapterError(
                    "Workbook has neither 'List of Trades' nor 'Strategy "
                    "Analysis' — not a MultiCharts Strategy Performance Report. "
                    f"Sheets found: {sorted(sheets)}"
                )
            settings = self._read_settings(wb) if "Settings" in sheets else {}
            trades, curve_seed = (
                self._read_trades(wb, settings)
                if "List of Trades" in sheets else ([], [])
            )
            summary = self._read_summary(wb, trades) if (
                "Strategy Analysis" in sheets
            ) else self._summary_from_trades(trades)
            self._fill_period(summary, settings, trades)
            equity_curve = self._build_equity(curve_seed, settings)
            symbol = settings.get("Symbol Name", "")
            compression = settings.get("Compression", "")
            session_name = " ".join(p for p in (symbol, compression) if p) \
                or "MultiCharts import"
            return NormalizedBacktestResult(
                summary=summary,
                trades=trades,
                equity_curve=equity_curve,
                source_app=self.app_id,
                session_name=session_name,
                settings=settings,
            )
        except BacktestAdapterError:
            raise
        except Exception as e:
            # v2.7 Task C (holistic-audit F7): defense-in-depth for the read
            # body. In the INSTALLED openpyxl the known adversarial shapes
            # (chartsheet under a data-sheet name, truncated sheet XML)
            # happen to raise at load_workbook and hit the open guard above
            # — but only row CONTENT parsing is deferred-by-contract, so a
            # version/shape that defers further would surface here and
            # escape the route's `except ValueError` as a 500. The wrap
            # also converts genuine adapter bugs into operator banners —
            # log loudly so the traceback isn't lost (observability rule).
            log.warning(
                "[multicharts] workbook read failed post-open: %s",
                e, exc_info=True,
            )
            raise BacktestAdapterError(
                "Failed reading the workbook (corrupt sheet data, unexpected "
                f"sheet type, or malformed cells): {type(e).__name__}: {e}"
            ) from e
        finally:
            wb.close()

    # ── Settings sheet ──────────────────────────────────────────────────────

    def _read_settings(self, wb) -> Dict[str, str]:
        """Key/value pairs from cols A/B, values stringified (JSON-safe)."""
        out: Dict[str, str] = {}
        for row in wb["Settings"].iter_rows(max_col=2, values_only=True):
            key, val = row[0], row[1] if len(row) > 1 else None
            if key is None or val is None:
                continue
            out[str(key).strip()] = _to_iso(val) if isinstance(val, datetime) \
                else str(val)
        return out

    # ── List of Trades sheet ────────────────────────────────────────────────

    def _read_trades(
        self, wb, settings: Dict[str, str]
    ) -> Tuple[List[BacktestTrade], List[Tuple[str, Optional[float], Optional[float]]]]:
        """Pair Entry/Exit rows into BacktestTrade list.

        Returns (trades, curve_seed) where curve_seed carries
        (exit_dt, cum_profit, drawdown) per closed trade for the equity
        curve (those columns live on the ENTRY row).
        """
        ws = wb["List of Trades"]
        rows = ws.iter_rows(values_only=True)

        header_map: Dict[str, int] = {}
        for row in rows:
            cells = [str(c).strip() if c is not None else "" for c in row]
            if "Trade #" in cells and "Type" in cells:
                header_map = {name: i for i, name in enumerate(cells) if name}
                break
        if not header_map:
            log.warning("[multicharts] no trade-list header found — 0 trades")
            return [], []

        def cell(row, name):
            i = header_map.get(name)
            return row[i] if i is not None and i < len(row) else None

        point_value = _to_float(settings.get("Point Value"))
        if point_value is None or point_value <= 0:
            log.warning(
                "[multicharts] Settings 'Point Value' missing/invalid — "
                "size_usdt falls back to Price x Contracts (point value 1.0)"
            )
            point_value = 1.0
        symbol = settings.get("Symbol Name", "")

        trades: List[BacktestTrade] = []
        curve_seed: List[Tuple[str, Optional[float], Optional[float]]] = []
        pending: Optional[Dict[str, Any]] = None
        skipped = 0

        for row in rows:  # continues from the row after the header
            typ = str(cell(row, "Type") or "").strip()
            if typ.startswith("Entry"):
                if pending is not None:
                    skipped += 1  # entry with no exit before the next entry
                pending = {
                    "side": "long" if "Long" in typ else "short",
                    "dt": _to_iso(cell(row, "Date") or cell(row, "Time")),
                    "price": _to_float(cell(row, "Price")) or 0.0,
                    "contracts": _to_float(cell(row, "Contracts")) or 0.0,
                    "pnl": _to_float(cell(row, "Profit ($)")),
                    "cum": _to_float(cell(row, "Cum. Profit ($)")),
                    "dd": _to_float(cell(row, "Drawdown ($)")),
                }
            elif typ.startswith("Exit"):
                if pending is None:
                    skipped += 1  # exit with no matching entry
                    continue
                exit_dt = _to_iso(cell(row, "Date") or cell(row, "Time"))
                signal = str(cell(row, "Signal") or "").strip()
                if "TP" in signal.upper():
                    exit_reason = "take_profit"
                elif "SL" in signal.upper():
                    exit_reason = "stop_loss"
                else:
                    exit_reason = signal.lower()
                trades.append(BacktestTrade(
                    symbol=symbol,
                    side=pending["side"],
                    entry_dt=pending["dt"],
                    exit_dt=exit_dt,
                    entry_price=pending["price"],
                    exit_price=_to_float(cell(row, "Price")) or 0.0,
                    size_usdt=pending["price"] * pending["contracts"] * point_value,
                    pnl_usdt=pending["pnl"] or 0.0,
                    exit_reason=exit_reason,
                ))
                curve_seed.append((exit_dt, pending["cum"], pending["dd"]))
                pending = None

        if pending is not None:
            skipped += 1  # trailing open position at export time
        if skipped:
            log.info("[multicharts] skipped %d unpaired trade rows", skipped)
        return trades, curve_seed

    # ── equity curve (from the trade list — the Graphs sheets are images) ──

    def _build_equity(
        self,
        curve_seed: List[Tuple[str, Optional[float], Optional[float]]],
        settings: Dict[str, str],
    ) -> List[EquityPoint]:
        initial = _to_float(settings.get("Initial Capital")) or 0.0
        points: List[EquityPoint] = []
        for dt, cum, dd in curve_seed:
            if cum is None:
                continue
            points.append(EquityPoint(
                dt=dt,
                equity=initial + cum,
                drawdown=abs(dd) if dd is not None else 0.0,
            ))
        return points

    # ── Strategy Analysis sheet ─────────────────────────────────────────────

    def _read_summary(self, wb, trades: List[BacktestTrade]) -> BacktestSummary:
        labels: Dict[str, Any] = {}
        for row in wb["Strategy Analysis"].iter_rows(max_col=2, values_only=True):
            if row[0] is not None and len(row) > 1 and row[1] is not None:
                labels[str(row[0]).strip()] = row[1]
        if not labels:
            # v2.7 Task C (holistic-audit F10): a layout variant that parses
            # ZERO pairs previously produced a silent, plausible-looking
            # summary. Tolerance contract says sparse ≠ loud, but zero
            # deserves a signal.
            log.warning(
                "[multicharts] Strategy Analysis parsed ZERO label/value "
                "pairs — summary falls back to trade-derived basics"
            )

        def num(label) -> Optional[float]:
            return _to_float(labels.get(label))

        fallback = self._summary_from_trades(trades)
        gross_profit = num("Gross Profit")
        gross_loss = num("Gross Loss")
        if gross_profit is not None and gross_loss:
            # MC reports Profit Factor signed-negative — recompute unsigned.
            profit_factor = gross_profit / abs(gross_loss)
        else:
            # F10 fallback symmetry: prefer the cell, else the
            # trades-derived PF (was hard-coded 0.0 one line after the
            # fallback had computed the true value).
            pf_cell = num("Profit Factor")
            profit_factor = abs(pf_cell) if pf_cell is not None \
                else fallback.profit_factor
        net_profit = num("Net Profit")
        win_rate = num("% Profitable")
        max_dd = num("Max Strategy Drawdown")
        max_dd_pct = num("Max Strategy Drawdown (%)")
        sharpe = num("Annualized Sharpe Ratio")
        if sharpe is None:
            sharpe = num("Sharpe Ratio")
        return BacktestSummary(
            net_profit=net_profit if net_profit is not None else fallback.net_profit,
            win_rate=win_rate if win_rate is not None else fallback.win_rate,
            profit_factor=profit_factor,
            max_drawdown=abs(max_dd) if max_dd is not None else 0.0,
            max_drawdown_pct=abs(max_dd_pct) if max_dd_pct is not None else 0.0,
            sharpe=sharpe if sharpe is not None else 0.0,
            # F7 int-cast guard: int(inf) raises OverflowError (escaped the
            # route's ValueError catch), int(nan) a cryptic ValueError —
            # and nan is TRUTHY so `or fallback` never engaged. Clean
            # fallback instead of a banner.
            total_trades=_to_int(num("Total # of Trades"),
                                 fallback.total_trades),
        )

    @staticmethod
    def _summary_from_trades(trades: List[BacktestTrade]) -> BacktestSummary:
        """Degraded lane: Strategy Analysis sheet absent — derive the basics."""
        if not trades:
            return BacktestSummary()
        wins = [t for t in trades if t.pnl_usdt > 0]
        gross_p = sum(t.pnl_usdt for t in wins)
        gross_l = sum(t.pnl_usdt for t in trades if t.pnl_usdt < 0)
        return BacktestSummary(
            net_profit=sum(t.pnl_usdt for t in trades),
            win_rate=len(wins) / len(trades),
            profit_factor=(gross_p / abs(gross_l)) if gross_l else 0.0,
            total_trades=len(trades),
        )

    @staticmethod
    def _fill_period(
        summary: BacktestSummary,
        settings: Dict[str, str],
        trades: List[BacktestTrade],
    ) -> None:
        summary.period_start = settings.get("Start Date") \
            or (trades[0].entry_dt if trades else "")
        summary.period_end = settings.get("End Date") \
            or (trades[-1].exit_dt if trades else "")
