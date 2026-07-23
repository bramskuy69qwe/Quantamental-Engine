"""
v2.7 Phase 2 — backtest-adapter framework + MultiCharts parser.

Covers (docs/design/v2.7_model_library_plan.md, Phase 2):
- registry resolution (mirrors core/adapters/registry coverage)
- happy path for BOTH fixture extensions (.xlsx and .xml — MC's .xml is
  the same OOXML zip; the adapter sniffs magic bytes, not extensions)
- the verified real-layout gotchas: header on row 3, signed-negative
  Max Strategy Drawdown / Profit Factor, "$50" Point Value string,
  trailing unpaired Entry (open position) skipped, EOD exit signal
- tolerance: reordered/renamed/missing columns; missing Strategy
  Analysis sheet (summary derived from trades); missing Settings sheet
  (point-value + initial-capital fallbacks)
- negative: garbage bytes / flat-XML raise BacktestAdapterError
- acceptance: asdict(result) feeds the Phase-1 wrappers unchanged
  (key parity with backtest_trades / backtest_equity)
- layout-stability check against the REAL @ES export (skipped when the
  sample isn't on this machine — it lives outside the repo)

Fixtures: tests/fixtures/multicharts/ (synthetic; regenerate with
make_fixture.py in the same dir — known values documented there).
"""
from __future__ import annotations

import os
from dataclasses import asdict
from io import BytesIO
from pathlib import Path

import pytest

from core.backtest_adapters import (
    BacktestAdapterError,
    get_backtest_adapter,
    list_backtest_adapters,
)
from core.database import DatabaseManager

FIXTURES = Path(__file__).parent / "fixtures" / "multicharts"
REAL_SAMPLE = (
    r"E:\Quantamental Models"
    r"\Backtesting Strategy Performance Report _ @ES - 1 Minute.xlsx"
)


def _fixture_bytes(name: str = "mc_synthetic.xlsx") -> bytes:
    return (FIXTURES / name).read_bytes()


def _parse(file_bytes: bytes, filename: str = "mc_synthetic.xlsx"):
    return get_backtest_adapter("multicharts").parse(file_bytes, filename)


# ── Registry ─────────────────────────────────────────────────────────────────

def test_registry_resolves_multicharts():
    adapter = get_backtest_adapter("multicharts")
    assert adapter.app_id == "multicharts"
    assert adapter.display_name == "MultiCharts"
    assert set(adapter.accepted_extensions) == {".xlsx", ".xml"}


def test_registry_unknown_app_raises_with_available_list():
    with pytest.raises(ValueError, match="multicharts"):
        get_backtest_adapter("tradestation")


def test_registry_listing_drives_upload_dropdown():
    rows = list_backtest_adapters()
    mc = next(r for r in rows if r["app_id"] == "multicharts")
    assert mc["display_name"] == "MultiCharts"
    assert ".xlsx" in mc["accepted_extensions"]


# ── Happy path (both extensions) ─────────────────────────────────────────────

def test_parse_xlsx_summary_trades_equity():
    result = _parse(_fixture_bytes())

    s = result.summary
    assert s.net_profit == 250.0
    assert s.win_rate == 0.5
    # Fixture's Profit Factor CELL is -1.9 (deliberately inconsistent):
    # 2.0 here proves the gross-recompute path, not abs(cell).
    assert s.profit_factor == 2.0
    assert s.max_drawdown == 120.0         # abs of MC's -120
    assert s.max_drawdown_pct == 0.0012    # abs of MC's -0.0012
    assert s.sharpe == 1.25                # Annualized preferred over plain
    assert s.total_trades == 10
    assert s.period_start.startswith("2026-06-24T09:30")
    assert s.period_end.startswith("2026-06-24T")

    trades = result.trades
    assert len(trades) == 10               # trailing unpaired Entry skipped
    t0 = trades[0]
    assert t0.symbol == "@ES"
    assert t0.side == "long"
    assert t0.entry_price == 5000.0
    assert t0.size_usdt == 5000.0 * 2 * 50.0   # Price x Contracts x PointValue
    assert t0.pnl_usdt == 100.0
    assert t0.exit_reason == "take_profit"
    assert t0.entry_dt.startswith("2026-06-24T09:30")
    # Both sides in both outcomes (fixture de-confounds side vs result).
    assert trades[1].side == "short" and trades[1].exit_reason == "take_profit"
    assert trades[5].side == "short" and trades[5].exit_reason == "stop_loss"
    assert trades[6].side == "long" and trades[6].exit_reason == "stop_loss"
    assert trades[9].exit_reason == "eod"  # non-TP/SL signal → raw lowercase
    assert all(t.r_multiple == 0.0 and t.regime_label == "" for t in trades)

    curve = result.equity_curve
    assert len(curve) == 10
    assert curve[0].equity == 100000.0 + 100.0   # InitialCapital + Cum. Profit
    assert curve[0].drawdown == 25.0             # abs of the -25 cell
    assert curve[-1].equity == 100000.0 + 250.0
    assert curve[0].dt.startswith("2026-06-24T09:35")

    assert result.source_app == "multicharts"
    assert result.session_name == "@ES 1 Minute"
    assert result.settings.get("Point Value") == "$50"


def test_parse_xml_extension_same_bytes_same_result():
    """MC's .xml export is the same OOXML zip — sniffed by magic bytes."""
    via_xlsx = _parse(_fixture_bytes("mc_synthetic.xlsx"), "r.xlsx")
    via_xml = _parse(_fixture_bytes("mc_synthetic.xml"), "r.xml")
    assert asdict(via_xml) == asdict(via_xlsx)


# ── Tolerance (sparse/reshaped files must not crash) ─────────────────────────

def _mutated_fixture(mutator) -> bytes:
    """Load the fixture, apply a mutation, return re-saved bytes."""
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(_fixture_bytes()))
    mutator(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tolerates_reordered_and_missing_columns():
    def reorder(wb):
        ws = wb["List of Trades"]
        # Rewrite the header with columns shuffled and Run-up/Drawdown gone;
        # move the data columns to match. Simplest: rebuild the sheet.
        rows = list(ws.iter_rows(values_only=True))
        header = list(rows[2])
        # 'Date' deliberately dropped, 'Time' kept — pins the Time fallback.
        keep = ["Type", "Signal", "Trade #", "Price", "Contracts",
                "Profit ($)", "Cum. Profit ($)", "Time"]
        idx = {name: header.index(name) for name in keep}
        del wb["List of Trades"]
        ws2 = wb.create_sheet("List of Trades")
        ws2.append(["List of Trades"])
        ws2.append([])
        ws2.append(keep)  # header on row 3, shuffled order, columns missing
        for row in rows[3:]:
            ws2.append([row[idx[name]] if idx[name] < len(row) else None
                        for name in keep])
    result = _parse(_mutated_fixture(reorder))
    assert len(result.trades) == 10
    assert result.trades[0].side == "long"
    assert result.trades[0].pnl_usdt == 100.0
    # 'Date' column absent → the 'Time' cell fills entry/exit datetimes.
    assert result.trades[0].entry_dt.startswith("2026-06-24T09:30")
    # Drawdown column absent → equity drawdown defaults to 0, no crash.
    assert len(result.equity_curve) == 10
    assert result.equity_curve[0].drawdown == 0.0


def test_tolerates_missing_strategy_analysis_sheet():
    def drop(wb):
        del wb["Strategy Analysis"]
    result = _parse(_mutated_fixture(drop))
    s = result.summary
    # Derived from the trade list instead: 5 x +100 / 5 x -50.
    assert s.net_profit == 250.0
    assert s.win_rate == 0.5
    assert s.profit_factor == 2.0
    assert s.total_trades == 10
    assert s.max_drawdown == 0.0  # not derivable — defaulted, not invented


def test_tolerates_missing_settings_sheet():
    def drop(wb):
        del wb["Settings"]
    result = _parse(_mutated_fixture(drop))
    # Point Value fallback 1.0 → notional = Price x Contracts.
    assert result.trades[0].size_usdt == 5000.0 * 2
    # Initial Capital fallback 0 → equity = Cum. Profit.
    assert result.equity_curve[0].equity == 100.0
    # Period falls back to first entry / last exit.
    assert result.summary.period_start.startswith("2026-06-24T09:30")
    assert result.session_name == "MultiCharts import"


def test_percent_string_cells_convert_to_fractions():
    """P2 audit MED-1: a future MC format exporting '% Profitable' as the
    string '37.2%' must land as 0.372, not a silent 100x drift."""
    def stringify_pct(wb):
        ws = wb["Strategy Analysis"]
        for row in ws.iter_rows():
            if row[0].value == "% Profitable":
                row[1].value = "37.2%"
    result = _parse(_mutated_fixture(stringify_pct))
    assert result.summary.win_rate == pytest.approx(0.372)


def test_zero_data_rows_yields_empty_trades_not_crash():
    """Tolerance lane: header present, no trade rows — trades/equity empty;
    the Strategy Analysis summary still reports its own totals (documented
    mismatch a Phase-3 consumer may surface, not a parse failure)."""
    def gut_rows(wb):
        ws = wb["List of Trades"]
        ws.delete_rows(4, ws.max_row)
    result = _parse(_mutated_fixture(gut_rows))
    assert result.trades == []
    assert result.equity_curve == []
    assert result.summary.total_trades == 10  # from the summary sheet


def test_unrecognizable_workbook_raises():
    """A valid workbook with NEITHER data sheet is not an MC report."""
    def gut(wb):
        del wb["List of Trades"]
        del wb["Strategy Analysis"]
    with pytest.raises(BacktestAdapterError, match="not a MultiCharts"):
        _parse(_mutated_fixture(gut))


# ── Negative (unrecognizable formats) ────────────────────────────────────────

def test_garbage_bytes_raise():
    with pytest.raises(BacktestAdapterError, match="unrecognized format"):
        _parse(b"\x00\x01\x02 definitely not a workbook", "junk.bin")


def test_flat_xml_raises_specific_error():
    flat = b"<?xml version='1.0'?><Workbook><Row/></Workbook>"
    with pytest.raises(BacktestAdapterError, match="flat-XML"):
        _parse(flat, "report.xml")


def test_pk_prefixed_garbage_raises_workbook_error():
    with pytest.raises(BacktestAdapterError, match="Could not open workbook"):
        _parse(b"PK\x03\x04 but not really a zip", "fake.xlsx")


# ── Acceptance: asdict() drops into the Phase-1 wrappers unchanged ──────────

def test_trade_and_equity_key_parity_with_db_columns():
    result = _parse(_fixture_bytes())
    trade_keys = set(asdict(result.trades[0]).keys())
    assert trade_keys == {
        "symbol", "side", "entry_dt", "exit_dt", "entry_price", "exit_price",
        "size_usdt", "r_multiple", "pnl_usdt", "regime_label", "exit_reason",
        "contracts",  # v3.0 P7 (G-M7): raw contract count, column added
    }
    equity_keys = set(asdict(result.equity_curve[0]).keys())
    assert equity_keys == {"dt", "equity", "drawdown"}


@pytest.mark.asyncio
async def test_end_to_end_import_via_phase1_wrappers(tmp_path):
    """Acceptance criterion: parsed result → create_model_backtest with
    asdict() only — no reshaping."""
    result = _parse(_fixture_bytes())
    mgr = DatabaseManager(path=str(tmp_path / "e2e.db"))
    await mgr.initialize()
    try:
        mid = await mgr.create_potential_model("MC model", "micro", "", {})
        sid = await mgr.create_model_backtest(mid, result.source_app, asdict(result))

        runs = await mgr.list_model_backtests(mid)
        assert len(runs) == 1
        assert runs[0]["status"] == "completed"
        assert runs[0]["summary"]["net_profit"] == 250.0
        assert runs[0]["summary"]["max_drawdown_pct"] == 0.0012
        assert runs[0]["date_from"].startswith("2026-06-24T09:30")

        trades = await mgr.get_backtest_trades(sid)
        assert len(trades) == 10
        assert trades[0]["symbol"] == "@ES"
        assert trades[0]["size_usdt"] == 500000.0
        equity = await mgr.get_backtest_equity(sid)
        assert len(equity) == 10
        assert equity[0]["drawdown"] == 25.0
    finally:
        await mgr._conn.close()


# ── Layout stability vs the REAL export (machine-local; skipped elsewhere) ──

@pytest.mark.skipif(
    not os.path.exists(REAL_SAMPLE),
    reason="real @ES sample lives outside the repo (operator machine only)",
)
def test_real_sample_parses_with_verified_shape():
    """§6-1 layout-stability check: the actual 1.17 MB @ES export parses,
    with the values inspected on 2026-07-16."""
    result = _parse(Path(REAL_SAMPLE).read_bytes(), os.path.basename(REAL_SAMPLE))
    s = result.summary
    assert s.total_trades == 129
    assert s.net_profit == -937.5
    assert s.max_drawdown == 1162.5            # abs of -1162.5
    assert 0 < s.win_rate < 1                  # 0.372…, a fraction
    assert s.profit_factor == pytest.approx(2512.5 / 3450.0)  # unsigned
    assert len(result.trades) > 100
    assert result.session_name == "@ES 1 Minute"
    assert all(t.size_usdt > 0 for t in result.trades)  # $50 point value found
