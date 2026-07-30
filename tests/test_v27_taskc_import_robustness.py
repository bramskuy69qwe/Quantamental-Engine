"""
v2.7 holistic-audit Task C — import robustness (F7/F8/F9/F10/F14).

Pins:
- F7: the parse body's broad exception wrap — chartsheet-under-a-data-
  sheet-name (AttributeError) and corrupt-sheet-XML (ParseError/zlib)
  now surface as BacktestAdapterError (the route's error banner), not
  500; int-cast lanes ('inf'/nan Total # of Trades) fall back cleanly.
- F8: the parse runs via asyncio.to_thread (off the trading loop).
- F9: initialize()'s boot sweep marks interrupted ('running') IMPORTED
  runs as failed (engine-run sessions untouched); the run list renders
  a visible FAILED badge.
- F10: a Strategy Analysis sheet parsing ZERO pairs warns + the summary
  falls back to trade-derived values incl. profit_factor (was 0.0).
- F14: double-submit guards (hx-disabled-elt) on upload + model form;
  per-adapter accept filter (data-ext); declared-size precheck rejects
  before the body lands in RAM.

Run: pytest tests/test_v27_taskc_import_robustness.py -v
"""
from __future__ import annotations

import logging
import os
import sys
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
import pytest_asyncio
from fastapi import UploadFile
from starlette.requests import Request

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import api.routes_models as rm
from core.backtest_adapters import BacktestAdapterError, get_backtest_adapter
from core.database import DatabaseManager

FIXTURES = Path(__file__).parent / "fixtures" / "multicharts"


def _parse(data: bytes, filename: str = "r.xlsx"):
    return get_backtest_adapter("multicharts").parse(data, filename)


def _mutated_fixture(mutator) -> bytes:
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO((FIXTURES / "mc_synthetic.xlsx").read_bytes()))
    mutator(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── F7: exception surface ────────────────────────────────────────────────────

def test_chartsheet_under_data_sheet_name_is_adapter_error():
    """A Chartsheet named like a data sheet raises AttributeError inside
    openpyxl — the CLASS pin: it must surface as BacktestAdapterError
    (the route's banner), never a bare exception → 500. In the installed
    openpyxl it happens to raise at load_workbook (the open guard);
    the F7 body-wrap covers versions/shapes where sheet parsing defers
    to iter_rows — either message satisfies the contract."""
    import openpyxl
    wb = openpyxl.Workbook()
    sa = wb.active
    sa.title = "Strategy Analysis"
    sa.append(["Net Profit", 100.0])
    wb.create_chartsheet("List of Trades")
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(BacktestAdapterError, match="workbook"):
        _parse(buf.getvalue())


def test_corrupt_sheet_xml_is_adapter_error_not_500():
    """Truncated worksheet members inside a valid zip container — same
    class pin as above: BacktestAdapterError, never a bare
    ParseError/zlib.error → 500 (whether it surfaces at open or, in
    deferring versions, at iter_rows via the F7 body-wrap)."""
    src = (FIXTURES / "mc_synthetic.xlsx").read_bytes()
    zin = zipfile.ZipFile(BytesIO(src))
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = b"<worksheet><broken"
            zout.writestr(item, data)
    with pytest.raises(BacktestAdapterError, match="workbook"):
        _parse(out.getvalue())


def test_body_wrap_converts_deferred_read_failures(monkeypatch, caplog):
    """The F7 BODY-wrap pin (audit fold): the two adversarial shapes hit
    the OPEN guard in the installed openpyxl, leaving the new wrap
    untested — this discriminates it by raising INSIDE the read body,
    matching the wrap's own message (the open guard says 'Could not
    open'). Also pins the observability fold: the traceback is logged."""
    from core.backtest_adapters.multicharts import MultiChartsAdapter

    def boom(self, wb, settings):
        raise RuntimeError("simulated deferred sheet failure")

    monkeypatch.setattr(MultiChartsAdapter, "_read_trades", boom)
    caplog.set_level(logging.WARNING, logger="backtest_adapters.multicharts")
    with pytest.raises(BacktestAdapterError, match="Failed reading the workbook"):
        _parse((FIXTURES / "mc_synthetic.xlsx").read_bytes())
    assert any("read failed post-open" in r.getMessage() for r in caplog.records)


def test_non_finite_total_trades_falls_back_cleanly():
    """int('inf' cell) raised OverflowError (escaped the route's
    ValueError catch); nan was TRUTHY so `or fallback` never engaged.
    Both now fall back to the trades-derived count."""
    for cell in ("inf", "nan"):
        def poison(wb, cell=cell):
            ws = wb["Strategy Analysis"]
            for row in ws.iter_rows():
                if row[0].value == "Total # of Trades":
                    row[1].value = cell
        result = _parse(_mutated_fixture(poison))
        assert result.summary.total_trades == 10, cell  # fixture trade count


# ── F10: empty-summary signal + fallback symmetry ───────────────────────────

def test_zero_label_summary_warns_and_falls_back(caplog):
    def blank_col_b(wb):
        ws = wb["Strategy Analysis"]
        for row in ws.iter_rows():
            if len(row) > 1:
                row[1].value = None
    caplog.set_level(logging.WARNING, logger="backtest_adapters.multicharts")
    result = _parse(_mutated_fixture(blank_col_b))
    s = result.summary
    assert s.net_profit == 250.0        # trades-derived
    assert s.profit_factor == 2.0       # F10: was hard-coded 0.0
    assert s.win_rate == 0.5
    assert s.total_trades == 10
    assert any("ZERO label/value" in r.getMessage() for r in caplog.records)


# ── F9: boot sweep + visible status ─────────────────────────────────────────

@pytest.mark.asyncio
async def test_boot_sweep_marks_interrupted_imported_runs_failed(tmp_path):
    db_path = str(tmp_path / "sweep.db")
    mgr = DatabaseManager(path=db_path)
    await mgr.initialize()
    mid = await mgr.create_potential_model("M", "both", "", {})
    sid = await mgr.create_model_backtest(mid, "multicharts", {
        "session_name": "r", "summary": {"net_profit": 1.0},
        "trades": [], "equity_curve": [],
    })
    eng = await mgr.create_backtest_session("engine", "macro", "", "", {})
    # Simulate process death mid-import / mid-run.
    await mgr._conn.execute(
        "UPDATE backtest_sessions SET status='running' WHERE id IN (?, ?)",
        (sid, eng))
    await mgr._conn.commit()
    await mgr._conn.close()

    mgr2 = DatabaseManager(path=db_path)
    await mgr2.initialize()
    try:
        assert (await mgr2.get_backtest_session(sid))["status"] == "failed"
        # Engine-run lane (model_id NULL) is deliberately untouched.
        assert (await mgr2.get_backtest_session(eng))["status"] == "running"
    finally:
        await mgr2._conn.close()


# ── F8 + F14: route + template pins ─────────────────────────────────────────

def test_parse_runs_off_the_event_loop():
    """Source pin, __file__-anchored (audit fold: was cwd-relative). The
    BEHAVIORAL pin lives in the happy-path test below (to_thread spy)."""
    src = (Path(__file__).parent.parent / "api" / "routes_models.py"
           ).read_text(encoding="utf-8")
    assert "asyncio.to_thread(adapter.parse" in src


@pytest_asyncio.fixture
async def mdb(monkeypatch):
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    database = DatabaseManager(path=tmp.name)
    await database.initialize()
    monkeypatch.setattr(rm, "db", database)
    yield database
    await database.close()
    try:
        os.unlink(tmp.name)
        for ext in ("-wal", "-shm"):
            if os.path.exists(tmp.name + ext):
                os.unlink(tmp.name + ext)
    except OSError:
        pass


def _req() -> Request:
    return Request({
        "type": "http", "http_version": "1.1", "method": "POST",
        "path": "/x", "raw_path": b"/x", "root_path": "",
        "scheme": "http", "query_string": b"", "headers": [],
        "client": ("test", 0), "server": ("test", 80),
    })


class _NoReadIO(BytesIO):
    """A body that PROVES the precheck fired before materialization."""
    def read(self, *a):
        raise AssertionError("body was read despite the size precheck")


@pytest.mark.asyncio
async def test_size_precheck_rejects_before_read(mdb):
    """F14: the parser-measured size rejects BEFORE the body is
    materialized — proven with a body whose read() raises (audit fold:
    the old assertion couldn't tell precheck from post-read)."""
    mid = await mdb.create_potential_model("M", "both", "", {})
    big = UploadFile(file=_NoReadIO(b"tiny"), filename="r.xlsx",
                     size=rm.MAX_UPLOAD_BYTES + 1)
    resp = await rm.upload_model_backtest(_req(), mid, file=big,
                                          app_id="multicharts")
    # Fragments slim-down (2026-07-30): the route is JSON-only — errors
    # carry real status codes instead of the 200-banner htmx contract.
    assert resp.status_code == 400
    assert "File too large" in resp.body.decode()
    assert await mdb.list_model_backtests(mid) == []


@pytest.mark.asyncio
async def test_upload_happy_path_still_works_through_to_thread(mdb, monkeypatch):
    """F8: the off-loop parse must not change the result — and the
    BEHAVIORAL pin (audit fold): a to_thread spy proves adapter.parse
    actually routed through it."""
    import asyncio as _aio
    seen = []
    orig = _aio.to_thread

    async def spy(fn, *a, **k):
        seen.append(getattr(fn, "__name__", ""))
        return await orig(fn, *a, **k)

    monkeypatch.setattr(rm.asyncio, "to_thread", spy)
    mid = await mdb.create_potential_model("M", "both", "", {})
    good = UploadFile(file=BytesIO((FIXTURES / "mc_synthetic.xlsx").read_bytes()),
                      filename="r.xlsx")
    resp = await rm.upload_model_backtest(_req(), mid, file=good,
                                          app_id="multicharts")
    assert resp.status_code == 200 and "text-red" not in resp.body.decode()
    # v3.0 P7: the verbatim capture is a SECOND openpyxl pass — both
    # CPU-bound calls must ride to_thread (same F8 rationale).
    assert "parse" in seen and "capture" in seen
    runs = await mdb.list_model_backtests(mid)
    assert len(runs) == 1 and runs[0]["summary"]["net_profit"] == 250.0
